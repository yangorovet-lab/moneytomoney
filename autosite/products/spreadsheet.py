"""Paid: Home Office Planner for Excel and Google Sheets."""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ACCENT = "1A6B54"
LIGHT = "E8F2ED"
INPUT = "FFF4D6"
GREY = "5D6470"

TITLE = Font(bold=True, size=16, color=ACCENT)
BOLD = Font(bold=True)
HEAD = Font(bold=True, color="FFFFFF")
MUTED = Font(italic=True, color=GREY, size=9)
HEAD_FILL = PatternFill("solid", fgColor=ACCENT)
LIGHT_FILL = PatternFill("solid", fgColor=LIGHT)
INPUT_FILL = PatternFill("solid", fgColor=INPUT)
THIN = Side(style="thin", color="C9CED3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

# Starting ratios of body height, based on average body proportions.
RATIOS = [
    ("Chair seat height", 0.25, "Floor to top of seat. Feet flat, thighs roughly level."),
    ("Sitting desk / keyboard height", 0.40, "Elbows at about 90 degrees, shoulders relaxed."),
    ("Top of monitor when sitting", 0.69, "Roughly your seated eye height. At or slightly below it."),
    ("Standing desk / keyboard height", 0.63, "Elbows at about 90 degrees while standing, in your usual shoes."),
    ("Top of monitor when standing", 0.93, "Roughly your standing eye height. At or slightly below it."),
]

GEAR = [
    ("Office chair", "Seating", "Must"),
    ("Desk or desk converter", "Desk", "Must"),
    ("Monitor arm or riser", "Screen", "Nice"),
    ("Laptop stand", "Screen", "Must"),
    ("External keyboard", "Input", "Must"),
    ("Mouse", "Input", "Must"),
    ("Desk lamp or monitor light bar", "Lighting", "Nice"),
    ("Footrest", "Seating", "Later"),
    ("Cable tray and ties", "Tidy", "Later"),
    ("Anti-fatigue mat", "Desk", "Later"),
    ("Headset", "Calls", "Nice"),
    ("Webcam", "Calls", "Later"),
]

AUDIT = [
    ("Chair", "My feet rest flat on the floor or a footrest."),
    ("Chair", "My hips are level with or slightly above my knees."),
    ("Chair", "There is a small gap between the seat edge and the back of my knees."),
    ("Chair", "My lower back is supported (chair lumbar or a cushion)."),
    ("Chair", "I lean back against the backrest instead of perching forward."),
    ("Desk", "My elbows are at about 90 degrees when I type."),
    ("Desk", "My shoulders stay relaxed, not raised, while I work."),
    ("Desk", "I can sit close to the desk without twisting."),
    ("Screen", "The top of my screen is at or slightly below eye level."),
    ("Screen", "My screen is about an arm's length away."),
    ("Screen", "My main screen is straight in front of me."),
    ("Screen", "If I use a laptop, it is raised and I use a separate keyboard."),
    ("Input", "My wrists stay straight while typing."),
    ("Input", "My mouse is right next to my keyboard."),
    ("Lighting", "There is no glare or reflection on my screen."),
    ("Lighting", "My window is beside the screen, not behind it or behind me."),
    ("Lighting", "My desk is lit well enough to read paper comfortably."),
    ("Habits", "I stand up or move at least once an hour."),
    ("Habits", "I take short eye breaks (20-20-20)."),
    ("Habits", "I change position during the day."),
]


def _widths(ws, widths):
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def _header_row(ws, row, labels):
    for i, text in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=i, value=text)
        cell.font, cell.fill, cell.border = HEAD, HEAD_FILL, BOX
        cell.alignment = Alignment(vertical="center")


def start_sheet(wb):
    ws = wb.active
    ws.title = "Start here"
    _widths(ws, {"A": 4, "B": 90})
    ws["B2"] = "Home Office Planner"
    ws["B2"].font = TITLE
    lines = [
        ("How to use", True),
        ("1. Ergonomic calculator: pick cm or inches, enter your height in the yellow cell. "
         "You get starting heights for your chair, desk and monitor.", False),
        ("2. Gear budget: set your total budget, adjust the item list, enter planned and actual prices. "
         "Totals update automatically.", False),
        ("3. Setup audit: answer Yes or No to 20 checks and get a score out of 100. "
         "Re-check after each change.", False),
        ("Yellow cells are for your input. Everything else is calculated.", False),
        ("Google Sheets: File > Import > Upload this file. Formulas and dropdowns carry over.", False),
        ("Good to know", True),
        ("The measurements are starting points based on average body proportions. Bodies differ, "
         "so fine-tune by comfort: relaxed shoulders, elbows around 90 degrees, feet supported, "
         "screen top at or slightly below eye level.", False),
        ("This planner is general information, not medical advice. If you have ongoing pain, "
         "numbness or tingling, see a doctor or physiotherapist.", False),
    ]
    row = 4
    for text, heading in lines:
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = Font(bold=True, color=ACCENT, size=12) if heading else Font(size=11)
        cell.alignment = WRAP
        row += 2 if heading else 1
        if not heading:
            ws.row_dimensions[row - 1].height = 32


def calculator_sheet(wb):
    ws = wb.create_sheet("Ergonomic calculator")
    _widths(ws, {"A": 34, "B": 14, "C": 14, "D": 64})
    ws["A1"] = "Ergonomic calculator"
    ws["A1"].font = TITLE

    ws["A3"], ws["A4"] = "Units", "Your height"
    for ref, value in (("B3", "cm"), ("B4", 175)):
        ws[ref] = value
        ws[ref].fill, ws[ref].border, ws[ref].font = INPUT_FILL, BOX, BOLD
    ws["C3"] = "cm or in"
    ws["C4"] = "without shoes"
    ws["C3"].font = ws["C4"].font = MUTED
    units = DataValidation(type="list", formula1='"cm,in"', allow_blank=False)
    height = DataValidation(type="decimal", operator="between", formula1="40", formula2="250")
    ws.add_data_validation(units)
    ws.add_data_validation(height)
    units.add("B3")
    height.add("B4")
    ws["A5"] = "Height in cm"
    ws["B5"] = '=IF(B3="in",B4*2.54,B4)'
    ws["B5"].number_format = "0"

    _header_row(ws, 7, ["Measurement", "cm", "inches", "What it means"])
    for i, (name, ratio, note) in enumerate(RATIOS, start=8):
        ws.cell(row=i, column=1, value=name).font = BOLD
        cm = ws.cell(row=i, column=2, value=f"=ROUND($B$5*{ratio},0)")
        inch = ws.cell(row=i, column=3, value=f"=ROUND(B{i}/2.54,1)")
        cm.number_format, inch.number_format = "0", "0.0"
        ws.cell(row=i, column=4, value=note).alignment = WRAP
        for col in range(1, 5):
            ws.cell(row=i, column=col).border = BOX
            if i % 2 == 0:
                ws.cell(row=i, column=col).fill = LIGHT_FILL
    row = 8 + len(RATIOS)
    ws.cell(row=row, column=1, value="Monitor distance").font = BOLD
    ws.cell(row=row, column=2, value="50-75")
    ws.cell(row=row, column=3, value="20-30")
    ws.cell(row=row, column=4, value="About an arm's length. Text should be readable without leaning in.")
    for col in range(1, 5):
        ws.cell(row=row, column=col).border = BOX
    ws.cell(row=row + 2, column=1,
            value="Starting estimates from average body proportions. Adjust by comfort; add shoe height when standing.").font = MUTED


def budget_sheet(wb):
    ws = wb.create_sheet("Gear budget")
    _widths(ws, {"A": 32, "B": 13, "C": 11, "D": 15, "E": 15, "F": 10, "G": 40})
    ws["A1"] = "Gear budget planner"
    ws["A1"].font = TITLE
    ws["A3"] = "Total budget"
    ws["B3"] = 500
    ws["B3"].fill, ws["B3"].border, ws["B3"].font = INPUT_FILL, BOX, BOLD
    ws["B3"].number_format = '"$"#,##0'

    first, last = 9, 9 + len(GEAR) + 7
    summary = [
        ("Planned total", f"=SUM(D{first}:D{last})"),
        ("Spent so far", f'=SUMIF(F{first}:F{last},"Yes",E{first}:E{last})'),
        ("Left in budget", "=B3-B5"),
    ]
    for r, (name, formula) in enumerate(summary, start=4):
        ws.cell(row=r, column=1, value=name).font = BOLD
        cell = ws.cell(row=r, column=2, value=formula)
        cell.number_format, cell.border = '"$"#,##0', BOX
    ws.conditional_formatting.add("B6", CellIsRule(operator="lessThan", formula=["0"],
                                                   font=Font(bold=True, color="B42318")))

    _header_row(ws, 8, ["Item", "Category", "Priority", "Planned price", "Actual price", "Bought?", "Notes / link"])
    priority = DataValidation(type="list", formula1='"Must,Nice,Later"')
    bought = DataValidation(type="list", formula1='"Yes,No"')
    ws.add_data_validation(priority)
    ws.add_data_validation(bought)
    rows = GEAR + [("", "", "")] * 7
    for r, (item, category, prio) in enumerate(rows, start=first):
        values = [item, category, prio, None, None, "No" if item else None, None]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = BOX
            if col in (4, 5):
                cell.number_format = '"$"#,##0'
                cell.fill = INPUT_FILL
        priority.add(f"C{r}")
        bought.add(f"F{r}")
    ws.conditional_formatting.add(f"A{first}:G{last}",
                                  FormulaRule(formula=[f'$F{first}="Yes"'], fill=LIGHT_FILL))
    ws.freeze_panes = "A9"
    ws.cell(row=last + 2, column=1,
            value="Tip: buy the Must items first. Free fixes (books under a laptop, a rolled towel for lumbar support) "
                  "can cover the rest while you save.").font = MUTED


def audit_sheet(wb):
    ws = wb.create_sheet("Setup audit")
    _widths(ws, {"A": 12, "B": 70, "C": 12})
    ws["A1"] = "Setup audit"
    ws["A1"].font = TITLE
    first, last = 7, 6 + len(AUDIT)
    ws["A3"], ws["B3"] = "Score", f'=ROUND(COUNTIF(C{first}:C{last},"Yes")/{len(AUDIT)}*100,0)&" / 100"'
    ws["A4"] = "Verdict"
    ws["B4"] = (f'=IF(COUNTIF(C{first}:C{last},"Yes")>=17,"Great setup. Keep the habits going.",'
                f'IF(COUNTIF(C{first}:C{last},"Yes")>=12,"Good start. Fix the No items one at a time.",'
                '"Start with the chair and screen rows: they make the biggest difference."))')
    ws["A3"].font = ws["A4"].font = ws["B3"].font = BOLD
    _header_row(ws, 6, ["Area", "Check", "Yes / No"])
    answers = DataValidation(type="list", formula1='"Yes,No"')
    ws.add_data_validation(answers)
    for r, (area, text) in enumerate(AUDIT, start=first):
        ws.cell(row=r, column=1, value=area)
        ws.cell(row=r, column=2, value=text).alignment = WRAP
        cell = ws.cell(row=r, column=3)
        cell.fill = INPUT_FILL
        answers.add(f"C{r}")
        for col in range(1, 4):
            ws.cell(row=r, column=col).border = BOX
    ws.conditional_formatting.add(f"C{first}:C{last}", CellIsRule(operator="equal", formula=['"Yes"'], fill=LIGHT_FILL))
    ws.conditional_formatting.add(f"C{first}:C{last}", CellIsRule(operator="equal", formula=['"No"'],
                                                                   font=Font(bold=True, color="B42318")))


def build(path) -> None:
    wb = Workbook()
    start_sheet(wb)
    calculator_sheet(wb)
    budget_sheet(wb)
    audit_sheet(wb)
    wb.save(path)


def build_preview(path) -> None:
    """A one-page PDF picture of the calculator with example results, used for the shop preview."""
    from . import pdfkit as k

    c = k.new_canvas(path, "Home Office Planner preview")
    y = k.header(c, "Home Office Planner", "Excel & Google Sheets  |  Calculator, gear budget, setup audit")
    y = k.section(c, y, "Ergonomic calculator (example: 175 cm / 5 ft 9 in)")
    for name, cm in estimates(175):
        k.label(c, k.MARGIN, y, name, size=11, color=k.INK)
        c.setFont("Helvetica-Bold", 11)
        c.setFillColor(k.ACCENT)
        c.drawRightString(k.PAGE_W - k.MARGIN, y, f"{cm} cm  /  {cm / 2.54:.1f} in")
        c.setStrokeColor(k.LINE)
        c.line(k.MARGIN, y - 7, k.PAGE_W - k.MARGIN, y - 7)
        y -= 26
    y -= 14
    y = k.section(c, y, "Gear budget")
    for item, category, prio in GEAR[:6]:
        k.label(c, k.MARGIN, y, item, size=10.5, color=k.INK)
        k.label(c, k.MARGIN + 250, y, category, size=10)
        k.label(c, k.MARGIN + 350, y, prio, size=10, color=k.ACCENT, bold=True)
        y -= 20
    y -= 14
    y = k.section(c, y, "Setup audit")
    for _, text in AUDIT[:6]:
        y = k.check_item(c, y, text)
    k.label(c, k.MARGIN, y - 10, "...and 14 more checks, scored out of 100.", size=10)
    k.footer(c)
    c.showPage()
    c.save()


def estimates(height_cm: float) -> list[tuple[str, int]]:
    """The calculator's results in Python, for showing an example on the product page."""
    return [(name, round(height_cm * ratio)) for name, ratio, _ in RATIOS]
